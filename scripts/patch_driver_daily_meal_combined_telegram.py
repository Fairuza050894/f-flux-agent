from pathlib import Path
from datetime import datetime
import ast
import shutil

ROOT = Path(__file__).resolve().parents[1]
CHECKER_PATH = ROOT / "skills" / "qa_automation" / "checker.py"

timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
backup_path = CHECKER_PATH.with_name(f"checker.py.backup_before_combined_telegram_{timestamp}")
shutil.copy2(CHECKER_PATH, backup_path)

text = CHECKER_PATH.read_text()


def replace_function(source_text, func_name, replacement):
    tree = ast.parse(source_text)
    lines = source_text.splitlines(keepends=True)

    target = None
    for node in tree.body:
        if isinstance(node, ast.FunctionDef) and node.name == func_name:
            target = node
            break

    if not target:
        return source_text, False

    start = target.lineno - 1
    end = target.end_lineno
    new_lines = lines[:start] + [replacement + "\n"] + lines[end:]
    return "".join(new_lines), True


combined_code = r'''
def qa_finalize_telegram_result(result, url=None, module_name=None, mode=None, environment="Sandbox"):
    """
    Finalizer shared untuk single suite dan combined suite.
    """

    if not isinstance(result, dict):
        return result

    try:
        if not result.get("execution_metadata"):
            result["execution_metadata"] = collect_qa_execution_metadata_basic(
                url=url or result.get("base_url") or "https://mobospace-sandbox.pancaran-group.co.id",
                module_name=module_name or result.get("module_name") or result.get("feature") or "Uang Makan Driver",
                mode=mode or result.get("mode") or "regression",
                environment=environment or result.get("environment") or "Sandbox",
                result=result,
            )

        result = apply_console_warning_policy(result)
        result = normalize_testing_summary_sections(result)

        result["testing_summary"] = inject_metadata_into_testing_summary(
            result.get("testing_summary", ""),
            result.get("execution_metadata", {}),
        )

        result = enrich_qa_report_artifacts_with_metadata(result)

    except Exception as exc:
        try:
            result["metadata_enrichment_error"] = str(exc)
        except Exception:
            pass

    return result


def _qa_click_label(page, label):
    scopes = [
        page.locator(".v-navigation-drawer, .v-navigation-drawer__content, .v-list, nav, aside").first,
        page.locator("body"),
    ]

    for scope in scopes:
        candidates = [
            scope.get_by_text(label, exact=True).first,
            page.locator(f".v-list-item:has-text('{label}')").last,
            page.locator(f"[role='button']:has-text('{label}')").last,
            scope.get_by_text(label, exact=False).first,
            page.get_by_text(label, exact=True).last,
            page.get_by_text(label, exact=False).last,
        ]

        for candidate in candidates:
            try:
                if candidate.count() <= 0:
                    continue

                candidate.scroll_into_view_if_needed(timeout=3000)
                candidate.wait_for(state="visible", timeout=4000)
                candidate.click(timeout=5000)

                try:
                    page.wait_for_load_state("networkidle", timeout=10000)
                except Exception:
                    pass

                page.wait_for_timeout(1500)
                return True

            except Exception:
                continue

    return False


def _qa_make_driver_meal_open_subpage(subfeature_name, labels):
    def open_subpage(page, test_cases, preferred_subpage="Monitoring"):
        clicked_label = None

        for label in labels:
            if _qa_click_label(page, label):
                clicked_label = label
                break

        if clicked_label:
            add_test_case(
                test_cases,
                scenario=f"Open Driver Meal {subfeature_name} subpage",
                precondition="Driver Meal menu group is visible",
                steps=f"Click {clicked_label} submenu",
                expected=f"Driver Meal {subfeature_name} subpage should be opened",
                actual=f"Opened submenu: {clicked_label}",
                status="PASS",
            )
        else:
            add_test_case(
                test_cases,
                scenario=f"Open Driver Meal {subfeature_name} subpage",
                precondition="Driver Meal menu group is visible",
                steps=f"Click {subfeature_name} submenu",
                expected=f"Driver Meal {subfeature_name} subpage should be opened",
                actual=f"{subfeature_name} submenu was not found",
                status="FAIL",
            )

    return open_subpage


def _qa_extract_status_from_summary(summary):
    import re

    match = re.search(r"Status:\s*(PASS|FAILED|NEED REVIEW)", summary or "")

    if match:
        return match.group(1)

    return "UNKNOWN"


def _qa_extract_count(summary, label):
    import re

    match = re.search(rf"- {label}:\s*(\d+)", summary or "")

    if match:
        return int(match.group(1))

    return 0


def _qa_extract_section(summary, header):
    if not summary:
        return "-"

    headers = [
        "Verified:",
        "Failed:",
        "Need Review:",
        "Bugs:",
        "Warnings / Known Issues:",
        "Evidence:",
        "Recommendation:",
        "Policy Note:",
    ]

    lines = summary.splitlines()
    start = -1

    for idx, line in enumerate(lines):
        if line.strip() == header:
            start = idx + 1
            break

    if start == -1:
        return "-"

    end = len(lines)

    for idx in range(start, len(lines)):
        if lines[idx].strip() in headers:
            end = idx
            break

    section_lines = [
        line for line in lines[start:end]
        if line.strip()
    ]

    return "\n".join(section_lines) if section_lines else "-"


def _qa_run_driver_meal_subfeature(name, open_func, validator_func, url, module_name, mode, environment):
    original_open = globals().get("open_uang_makan_driver_subpage")
    original_monitoring_validator = globals().get("validate_uang_makan_driver_monitoring_page")
    original_page_validator = globals().get("validate_uang_makan_driver_page")

    try:
        globals()["open_uang_makan_driver_subpage"] = open_func
        globals()["validate_uang_makan_driver_monitoring_page"] = validator_func
        globals()["validate_uang_makan_driver_page"] = validator_func

        result = _perform_audit_for_telegram_original(
            url,
            module_name,
            mode,
        )

        result = qa_finalize_telegram_result(
            result,
            url=url,
            module_name=f"Driver Daily Meal - {name}",
            mode=mode,
            environment=environment,
        )

        return result

    finally:
        if original_open:
            globals()["open_uang_makan_driver_subpage"] = original_open

        if original_monitoring_validator:
            globals()["validate_uang_makan_driver_monitoring_page"] = original_monitoring_validator

        if original_page_validator:
            globals()["validate_uang_makan_driver_page"] = original_page_validator


def _qa_write_driver_daily_meal_combined_spreadsheet(path, metadata, sub_results, overall_status, total_counts):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except Exception:
        return None

    wb = Workbook()

    ws = wb.active
    ws.title = "Combined Summary"

    ws.append(["Driver Daily Meal Combined Suite"])
    ws.append([])
    ws.append(["Overall Status", overall_status])
    ws.append(["Passed", total_counts.get("passed", 0)])
    ws.append(["Failed", total_counts.get("failed", 0)])
    ws.append(["Need Review", total_counts.get("need_review", 0)])
    ws.append(["Skipped", total_counts.get("skipped", 0)])
    ws.append(["Bugs Found", total_counts.get("bugs", 0)])
    ws.append(["Non-blocking Warnings", total_counts.get("warnings", 0)])

    ws2 = wb.create_sheet("Subfeatures")
    ws2.append(["Subfeature", "Status", "Passed", "Failed", "Need Review", "Skipped", "Bugs", "Warnings"])

    for item in sub_results:
        counts = item["counts"]
        ws2.append([
            item["name"],
            item["status"],
            counts["passed"],
            counts["failed"],
            counts["need_review"],
            counts["skipped"],
            counts["bugs"],
            counts["warnings"],
        ])

    ws3 = wb.create_sheet("Artifacts")
    ws3.append(["Subfeature", "Screenshot", "Report", "Spreadsheet"])

    for item in sub_results:
        result = item["result"]
        ws3.append([
            item["name"],
            result.get("screenshot_path", "-"),
            result.get("report_path", "-"),
            result.get("spreadsheet_path", "-"),
        ])

    ws4 = wb.create_sheet("Execution Metadata")
    ws4.append(["Field", "Value"])

    for key, value in (metadata or {}).items():
        ws4.append([key, str(value)])

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(horizontal="center")

        for column_cells in sheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter

            for cell in column_cells:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))

            sheet.column_dimensions[column_letter].width = min(max_length + 4, 90)

    wb.save(path)
    return str(path)


def perform_driver_daily_meal_combined_suite(url, module_name="Uang Makan Driver", mode="regression", environment="Sandbox"):
    """
    Combined suite untuk:
    - Monitoring
    - Exclude
    - History / Inquiry
    """

    from pathlib import Path
    from datetime import datetime

    suites = [
        {
            "name": "Monitoring",
            "open": _qa_make_driver_meal_open_subpage("Monitoring", ["Monitoring"]),
            "validator": validate_uang_makan_driver_monitoring_page,
        },
        {
            "name": "Exclude",
            "open": _qa_make_driver_meal_open_subpage("Exclude", ["Exclude", "Pengecualian", "Exception", "Exceptions"]),
            "validator": validate_driver_meal_exclude_page,
        },
        {
            "name": "History / Inquiry",
            "open": _qa_make_driver_meal_open_subpage("History / Inquiry", ["History", "Inquiry", "Riwayat"]),
            "validator": validate_driver_meal_history_page,
        },
    ]

    sub_results = []

    for suite in suites:
        result = _qa_run_driver_meal_subfeature(
            suite["name"],
            suite["open"],
            suite["validator"],
            url,
            module_name,
            mode,
            environment,
        )

        summary = result.get("testing_summary") or ""

        status = _qa_extract_status_from_summary(summary)

        counts = {
            "passed": _qa_extract_count(summary, "Passed"),
            "failed": _qa_extract_count(summary, "Failed"),
            "need_review": _qa_extract_count(summary, "Need Review"),
            "skipped": _qa_extract_count(summary, "Skipped"),
            "bugs": _qa_extract_count(summary, "Bugs Found"),
            "warnings": _qa_extract_count(summary, "Non-blocking Warnings"),
        }

        sub_results.append(
            {
                "name": suite["name"],
                "status": status,
                "counts": counts,
                "result": result,
            }
        )

    has_failed = any(item["status"] == "FAILED" for item in sub_results)
    has_need_review = any(item["status"] == "NEED REVIEW" for item in sub_results)

    if has_failed:
        overall_status = "FAILED"
    elif has_need_review:
        overall_status = "NEED REVIEW"
    else:
        overall_status = "PASS"

    total_counts = {
        "passed": sum(item["counts"]["passed"] for item in sub_results),
        "failed": sum(item["counts"]["failed"] for item in sub_results),
        "need_review": sum(item["counts"]["need_review"] for item in sub_results),
        "skipped": sum(item["counts"]["skipped"] for item in sub_results),
        "bugs": sum(item["counts"]["bugs"] for item in sub_results),
        "warnings": sum(item["counts"]["warnings"] for item in sub_results),
    }

    run_id = datetime.now().strftime("%Y%m%d_%H%M%S")
    base_dir = Path(__file__).resolve().parent / "artifacts"
    report_dir = base_dir / "reports"
    log_dir = base_dir / "logs"
    spreadsheet_dir = base_dir / "spreadsheets"

    report_dir.mkdir(parents=True, exist_ok=True)
    log_dir.mkdir(parents=True, exist_ok=True)
    spreadsheet_dir.mkdir(parents=True, exist_ok=True)

    metadata = collect_qa_execution_metadata_basic(
        url=url,
        module_name="Driver Daily Meal",
        mode=mode,
        environment=environment,
        result={"overall_status": overall_status},
    )

    subfeature_lines = "\n".join(
        f"- {item['name']}: {item['status']}"
        for item in sub_results
    )

    artifact_lines = []

    for item in sub_results:
        result = item["result"]
        artifact_lines.append(
            f"{item['name']}:\n"
            f"- Screenshot: {result.get('screenshot_path', '-')}\n"
            f"- Report: {result.get('report_path', '-')}\n"
            f"- Spreadsheet: {result.get('spreadsheet_path', '-')}"
        )

    warnings = []

    for item in sub_results:
        section = _qa_extract_section(item["result"].get("testing_summary", ""), "Warnings / Known Issues:")
        if section and section != "-":
            warnings.append(f"{item['name']}:\n{section}")

    warning_text = "\n\n".join(warnings) if warnings else "-"

    need_reviews = []

    for item in sub_results:
        section = _qa_extract_section(item["result"].get("testing_summary", ""), "Need Review:")
        if section and section != "-":
            need_reviews.append(f"{item['name']}:\n{section}")

    need_review_text = "\n\n".join(need_reviews) if need_reviews else "-"

    failed_items = []

    for item in sub_results:
        section = _qa_extract_section(item["result"].get("testing_summary", ""), "Failed:")
        if section and section != "-":
            failed_items.append(f"{item['name']}:\n{section}")

    failed_text = "\n\n".join(failed_items) if failed_items else "-"

    testing_summary = (
        "✅ QA E2E Regression Completed\n\n"
        "Module: Driver Daily Meal\n"
        f"Mode: {mode}\n"
        f"Environment: {environment}\n"
        f"Status: {overall_status}\n\n"
        "Execution Info:\n"
        f"- Execution ID: {metadata.get('execution_id', '-')}\n"
        f"- Executed At: {metadata.get('executed_at', '-')}\n"
        f"- Executed By: {metadata.get('executed_by', '-')}\n"
        "- Feature: Driver Daily Meal\n"
        f"- Suite / Mode: {mode}\n"
        f"- Environment: {environment}\n"
        f"- Base URL: {url}\n\n"
        "Runtime Environment:\n"
        f"- OS: {metadata.get('os', '-')}\n"
        f"- Machine: {metadata.get('machine', '-')}\n"
        f"- Browser: {metadata.get('browser', '-')}\n"
        f"- Browser Version: {metadata.get('browser_version', '-')}\n"
        f"- Device Profile: {metadata.get('device_profile', '-')}\n"
        f"- Viewport: {metadata.get('viewport', '-')}\n"
        f"- Automation Tool: {metadata.get('automation_tool', '-')}\n"
        f"- Runtime: {metadata.get('runtime', '-')}\n"
        f"- Python Version: {metadata.get('python_version', '-')}\n\n"
        "Summary:\n"
        f"- Passed: {total_counts['passed']}\n"
        f"- Failed: {total_counts['failed']}\n"
        f"- Need Review: {total_counts['need_review']}\n"
        f"- Skipped: {total_counts['skipped']}\n"
        f"- Bugs Found: {total_counts['bugs']}\n"
        f"- Non-blocking Warnings: {total_counts['warnings']}\n\n"
        "Subfeature Results:\n"
        f"{subfeature_lines}\n\n"
        "Failed:\n"
        f"{failed_text}\n\n"
        "Need Review:\n"
        f"{need_review_text}\n\n"
        "Warnings / Known Issues:\n"
        f"{warning_text}\n\n"
        "Evidence:\n"
        + "\n\n".join(artifact_lines)
        + "\n\nRecommendation:\n"
        + (
            "Review non-blocking warnings with frontend developer before release."
            if overall_status == "NEED REVIEW"
            else "No blocking issue found from automation result. Continue with manual business verification if needed."
        )
    )

    report_path = report_dir / f"qa_documentation_driver_daily_meal_combined_{run_id}.md"
    error_log_path = log_dir / f"qa_error_log_driver_daily_meal_combined_{run_id}.md"
    spreadsheet_path = spreadsheet_dir / f"qa_report_driver_daily_meal_combined_{run_id}.xlsx"

    documentation_report = (
        "# QA Documentation - Driver Daily Meal Combined Suite\n\n"
        + testing_summary
        + "\n\n---\n\n"
        + "## Subfeature Detail\n\n"
    )

    for item in sub_results:
        documentation_report += (
            f"### {item['name']}\n\n"
            + (item["result"].get("testing_summary") or "-")
            + "\n\n---\n\n"
        )

    error_log_report = (
        "# QA Error / Bug Log - Driver Daily Meal Combined Suite\n\n"
        f"Environment: {environment}\n"
        f"Mode: {mode}\n"
        f"Status: {overall_status}\n"
        f"Execution Time: {run_id}\n\n"
        "---\n\n"
        "## Failed Items\n\n"
        f"{failed_text}\n\n"
        "## Need Review Items\n\n"
        f"{need_review_text}\n\n"
        "## Warnings / Known Issues\n\n"
        f"{warning_text}\n"
    )

    report_path.write_text(documentation_report, encoding="utf-8")
    error_log_path.write_text(error_log_report, encoding="utf-8")

    _qa_write_driver_daily_meal_combined_spreadsheet(
        spreadsheet_path,
        metadata,
        sub_results,
        overall_status,
        total_counts,
    )

    screenshot_path = None

    for item in sub_results:
        candidate = item["result"].get("screenshot_path")
        if candidate:
            screenshot_path = candidate
            break

    combined_result = {
        "module_name": "Driver Daily Meal",
        "mode": mode,
        "environment": environment,
        "overall_status": overall_status,
        "execution_metadata": metadata,
        "testing_summary": testing_summary,
        "documentation_report": documentation_report,
        "error_log_report": error_log_report,
        "report_path": str(report_path),
        "documentation_path": str(report_path),
        "error_log_path": str(error_log_path),
        "spreadsheet_path": str(spreadsheet_path),
        "screenshot_path": screenshot_path,
        "sub_results": sub_results,
        "metadata_markdown_enriched": True,
        "metadata_excel_enriched": True,
    }

    return combined_result


def should_run_driver_daily_meal_combined(module_name, mode):
    import os

    enabled = os.getenv("QA_DRIVER_DAILY_MEAL_COMBINED", "true").lower() in ["1", "true", "yes", "on"]

    if not enabled:
        return False

    module_text = str(module_name or "").lower()
    mode_text = str(mode or "").lower()

    module_match = (
        "uang makan driver" in module_text
        or "driver daily meal" in module_text
        or "driver meal" in module_text
        or "meal allowance" in module_text
    )

    mode_match = mode_text in ["regression", "full", "e2e"]

    return module_match and mode_match
'''

if "def perform_driver_daily_meal_combined_suite(" not in text:
    marker = "def perform_audit_for_telegram"
    idx = text.find(marker)

    if idx == -1:
        raise SystemExit("perform_audit_for_telegram marker not found")

    text = text[:idx] + combined_code + "\n\n" + text[idx:]
    print("Inserted Driver Daily Meal combined suite functions")
else:
    print("Combined suite functions already exist")


wrapper_function = r'''def perform_audit_for_telegram(*args, **kwargs):
    """
    Wrapper final:
    - metadata masuk ke Telegram summary, Markdown, Excel
    - known Vue warning tidak menjadi blocking FAIL
    - Driver Daily Meal regression menjalankan Monitoring + Exclude + History
    """

    try:
        url = (
            kwargs.get("url")
            or kwargs.get("base_url")
            or (args[0] if len(args) > 0 else None)
            or "https://mobospace-sandbox.pancaran-group.co.id"
        )

        module_name = (
            kwargs.get("module_name")
            or (args[1] if len(args) > 1 else None)
            or "Uang Makan Driver"
        )

        mode = (
            kwargs.get("mode")
            or (args[2] if len(args) > 2 else None)
            or "regression"
        )

        environment = (
            kwargs.get("environment")
            or "Sandbox"
        )

        if should_run_driver_daily_meal_combined(module_name, mode):
            return perform_driver_daily_meal_combined_suite(
                url=url,
                module_name=module_name,
                mode=mode,
                environment=environment,
            )

        result = _perform_audit_for_telegram_original(*args, **kwargs)

        result = qa_finalize_telegram_result(
            result,
            url=url,
            module_name=module_name,
            mode=mode,
            environment=environment,
        )

        return result

    except Exception as exc:
        result = _perform_audit_for_telegram_original(*args, **kwargs)

        try:
            result["metadata_enrichment_error"] = str(exc)
        except Exception:
            pass

        return result
'''

text, wrapper_replaced = replace_function(text, "perform_audit_for_telegram", wrapper_function)

if not wrapper_replaced:
    raise SystemExit("perform_audit_for_telegram wrapper not found")

CHECKER_PATH.write_text(text)

print("Patched perform_audit_for_telegram for Driver Daily Meal combined suite")
print(f"Backup created: {backup_path}")
