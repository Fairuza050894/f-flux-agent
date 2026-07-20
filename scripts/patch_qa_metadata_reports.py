from pathlib import Path
from datetime import datetime
import re
import shutil

ROOT = Path(__file__).resolve().parents[1]
CHECKER_PATH = ROOT / "skills" / "qa_automation" / "checker.py"

if not CHECKER_PATH.exists():
    raise SystemExit(f"checker.py not found: {CHECKER_PATH}")

timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
backup_path = CHECKER_PATH.with_name(f"checker.py.backup_before_metadata_reports_{timestamp}")
shutil.copy2(CHECKER_PATH, backup_path)

text = CHECKER_PATH.read_text()

if "def collect_qa_execution_metadata_basic(" not in text:
    raise SystemExit(
        "Basic QA metadata helper not found. Jalankan step basic metadata dulu sebelum metadata Markdown/Excel."
    )

helper_code = r'''
def qa_metadata_markdown_block(metadata):
    if not metadata:
        return ""

    rows = [
        ("Execution ID", metadata.get("execution_id", "-")),
        ("Executed At", metadata.get("executed_at", "-")),
        ("Executed By", metadata.get("executed_by", "-")),
        ("Feature", metadata.get("feature", "-")),
        ("Suite / Mode", metadata.get("suite", "-")),
        ("Environment", metadata.get("environment", "-")),
        ("Base URL", metadata.get("base_url", "-")),
        ("OS", metadata.get("os", "-")),
        ("Machine", metadata.get("machine", "-")),
        ("Browser", metadata.get("browser", "-")),
        ("Browser Version", metadata.get("browser_version", "-")),
        ("Device Profile", metadata.get("device_profile", "-")),
        ("Viewport", metadata.get("viewport", "-")),
        ("Automation Tool", metadata.get("automation_tool", "-")),
        ("Runtime", metadata.get("runtime", "-")),
        ("Python Version", metadata.get("python_version", "-")),
    ]

    lines = [
        "## QA Execution Metadata",
        "",
        "| Field | Value |",
        "|---|---|",
    ]

    for key, value in rows:
        safe_value = str(value).replace("|", "\\|")
        lines.append(f"| {key} | {safe_value} |")

    return "\n".join(lines)


def inject_metadata_into_markdown_report(report_text, metadata):
    report_text = report_text or ""
    metadata_block = qa_metadata_markdown_block(metadata)

    if not metadata_block:
        return report_text

    if "## QA Execution Metadata" in report_text:
        return report_text

    stripped = report_text.lstrip()

    if stripped.startswith("#"):
        lines = report_text.splitlines()
        if lines:
            first_line = lines[0]
            rest = "\n".join(lines[1:]).strip()
            if rest:
                return first_line + "\n\n" + metadata_block + "\n\n" + rest
            return first_line + "\n\n" + metadata_block

    return metadata_block + "\n\n" + report_text


def update_spreadsheet_with_metadata(spreadsheet_path, metadata):
    if not spreadsheet_path or not metadata:
        return spreadsheet_path

    try:
        from pathlib import Path as _Path
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except Exception:
        return spreadsheet_path

    path = _Path(spreadsheet_path)

    if not path.exists():
        return spreadsheet_path

    try:
        wb = load_workbook(path)

        sheet_name = "Execution Metadata"

        if sheet_name in wb.sheetnames:
            del wb[sheet_name]

        ws = wb.create_sheet(sheet_name, 0)

        rows = [
            ("Execution ID", metadata.get("execution_id", "-")),
            ("Executed At", metadata.get("executed_at", "-")),
            ("Executed By", metadata.get("executed_by", "-")),
            ("Feature", metadata.get("feature", "-")),
            ("Suite / Mode", metadata.get("suite", "-")),
            ("Environment", metadata.get("environment", "-")),
            ("Base URL", metadata.get("base_url", "-")),
            ("OS", metadata.get("os", "-")),
            ("Machine", metadata.get("machine", "-")),
            ("Browser", metadata.get("browser", "-")),
            ("Browser Version", metadata.get("browser_version", "-")),
            ("Device Profile", metadata.get("device_profile", "-")),
            ("Viewport", metadata.get("viewport", "-")),
            ("Automation Tool", metadata.get("automation_tool", "-")),
            ("Runtime", metadata.get("runtime", "-")),
            ("Python Version", metadata.get("python_version", "-")),
        ]

        ws.append(["Field", "Value"])

        for row in rows:
            ws.append(list(row))

        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row in ws.iter_rows(min_row=2):
            row[0].font = Font(bold=True)
            row[0].alignment = Alignment(vertical="top")
            row[1].alignment = Alignment(wrap_text=True, vertical="top")

        for column_cells in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)

            for cell in column_cells:
                value = cell.value
                if value:
                    max_length = max(max_length, len(str(value)))

            ws.column_dimensions[column_letter].width = min(max_length + 4, 80)

        wb.save(path)
    except Exception:
        return spreadsheet_path

    return spreadsheet_path


def enrich_qa_report_artifacts_with_metadata(result):
    result = result or {}
    metadata = result.get("execution_metadata") or {}

    if not metadata:
        return result

    documentation_report = result.get("documentation_report")

    if documentation_report is not None:
        enriched_report = inject_metadata_into_markdown_report(documentation_report, metadata)
        result["documentation_report"] = enriched_report

        report_path = (
            result.get("report_path")
            or result.get("documentation_path")
        )

        if report_path:
            try:
                from pathlib import Path as _Path
                _Path(report_path).write_text(enriched_report, encoding="utf-8")
            except Exception:
                pass

    spreadsheet_path = result.get("spreadsheet_path")

    if spreadsheet_path:
        update_spreadsheet_with_metadata(spreadsheet_path, metadata)

    return result

'''

if "def enrich_qa_report_artifacts_with_metadata(" not in text:
    marker = "def build_testing_summary"
    index = text.find(marker)

    if index == -1:
        raise SystemExit("Could not find def build_testing_summary marker")

    text = text[:index] + helper_code + "\n\n" + text[index:]
    print("Inserted metadata report helper functions")
else:
    print("Metadata report helper functions already exist")

if "enrich_qa_report_artifacts_with_metadata(result)" not in text:
    target = '''result["testing_summary"] = inject_metadata_into_testing_summary(
        result.get("testing_summary", ""),
        result.get("execution_metadata", {}),
    )'''

    if target not in text:
        raise SystemExit(
            "Could not find metadata testing_summary injection block. "
            "Kirim grep hasil perform_audit_for_telegram kalau ini terjadi."
        )

    replacement = target + '''

    result = enrich_qa_report_artifacts_with_metadata(result)'''

    text = text.replace(target, replacement)
    print("Inserted artifact metadata enrichment call")
else:
    print("Artifact metadata enrichment call already exists")

CHECKER_PATH.write_text(text)

print(f"Patch completed.")
print(f"Backup created: {backup_path}")
